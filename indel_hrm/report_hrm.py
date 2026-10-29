"""Outputs: XLSX workbook (per-pair separation, trees, primers) and per-sample images.

Generic: uses an optional code->name map (metadata CSV with columns 'code,name');
without it, uses the sample codes themselves.
"""
import csv
import json
import itertools
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .genotype import grid_of, pair_separation

HRMLAB = {"A": "LONG", "B": "SHORT", "AB": "HETEROZYGOUS", "N": "N (presence/absence)"}
_HF = PatternFill("solid", fgColor="1F4E78")
_HFONT = Font(color="FFFFFF", bold=True)
_THIN = Side(style="thin", color="CCCCCC")
_BD = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_PALETTE = ["#1f4e78", "#c55a11", "#2e7d32", "#7030a0", "#b00020", "#008b8b",
            "#8b4513", "#444444", "#d4a017", "#1565c0"]


def load_names(metadata_csv):
    names = {}
    if metadata_csv and Path(metadata_csv).exists():
        with open(metadata_csv, newline="") as fh:
            for row in csv.DictReader(fh):
                key = (row.get("code") or row.get("codigo") or "").strip()
                val = (row.get("name") or row.get("nome") or "").strip()
                if key:
                    names[key] = val or key
    return names


def _name(names, code):
    return names.get(code, code)


def _loc(r):
    return f"{r['chrom'].split('_chr')[-1]}:{r['pos']}"


def _hdr(ws):
    for c in ws[ws.max_row]:
        c.fill = _HF
        c.font = _HFONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _tree_markers(t):
    keys = {}

    def rec(entry):
        _, st = entry
        if st.get("leaf"):
            return
        keys[(st["node"]["chrom"], st["node"]["pos"])] = st["node"]
        for _, sub in st["groups"].items():
            rec(sub)

    keys[(t["node"]["chrom"], t["node"]["pos"])] = t["node"]
    for _, sub in t["groups"].items():
        rec(sub)
    return keys


def _save_images(result, cfg, names, img_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    img_dir = Path(img_dir)
    img_dir.mkdir(parents=True, exist_ok=True)
    grid = grid_of(cfg)
    samples = result["samples"]
    color = {s: _PALETTE[i % len(_PALETTE)] for i, s in enumerate(samples)}

    used = {}
    for t in result.get("trees", []):
        used.update(_tree_markers(t))
    if not used and result.get("usable"):  # panel: use the panel markers
        used = {(r["chrom"], r["pos"]): r for r in result["usable"][:12]}

    for key, r in used.items():
        fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4))
        ref = None
        for s in samples:
            st = r["labels"].get(s, "miss")
            if st not in r["curves"]:        # miss or N (no predictable curve)
                continue
            c = r["curves"][st]
            if ref is None:
                ref, refname = c, _name(names, s)
            a1.plot(grid, c * 100, color=color[s], label=f"{_name(names, s)} [{HRMLAB[st]}]")
            a2.plot(grid, (c - ref) * 100, color=color[s])
        a1.set_xlabel("Temperature (°C)")
        a1.set_ylabel("Normalized fluorescence (%)")
        a1.set_title("Per-sample predicted curves (by genotype)")
        a1.legend(fontsize=7)
        a1.grid(alpha=.3)
        a2.axhline(0, color="k", lw=.6)
        a2.set_xlabel("Temperature (°C)")
        a2.set_ylabel(f"Difference vs {refname} (%)")
        a2.set_title("Difference plot (predicted, uncalibrated)")
        a2.grid(alpha=.3)
        fig.suptitle(f"{_loc(r)} indel {r['indel']} bp — predicted separation by genotype")
        fig.tight_layout()
        fig.savefig(img_dir / f"{_loc(r).replace(':', '_')}_indel{r['indel']}.png", dpi=110)
        plt.close(fig)
    return img_dir


def _klass(cfg, peak):
    return "Easy" if peak >= cfg.hrm_easy else ("Likely" if peak >= cfg.hrm_likely else "Difficult")


def _insilico_str(r):
    ins = r.get("insilico")
    if not ins:
        return "n/a"
    per = ", ".join(f"{k}:{v}" for k, v in ins["per_hap"].items())
    return f"{per} ({'specific' if ins['specific'] else 'CHECK'})"


def write_report(result, cfg, out_dir, names=None):
    names = names or {}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    samples = result["samples"]
    wb = Workbook()

    # per-pair separation
    ws = wb.active
    ws.title = "Separation_per_pair"
    ws.append(["Separation between samples in HRM — per pair, using each sample's real GENOTYPE"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Pair", "Best marker", "Indel (bp)", "Predicted diff-plot peak (%, uncalibrated)",
               "Class", "Genotypes (X/Y)", "Evidence"])
    _hdr(ws)
    per_pair = result.get("per_pair", {})
    for (x, y) in itertools.combinations(samples, 2):
        rows = per_pair.get((x, y), [])
        lab = f"{_name(names, x)} × {_name(names, y)}"
        if not rows:
            ws.append([lab, "— none —", "", "", "", "", ""])
            continue
        for d in rows[:3]:
            ws.append([lab, _loc(d["marker"]), d["marker"]["indel"], d["peak"],
                       _klass(cfg, d["peak"]), f"{d['sx']}/{d['sy']}",
                       "STRONG" if d["forte"] else "uses N (confirm)"])
            for c in ws[ws.max_row]:
                c.border = _BD
        ws.append([])
    for col, w in {"A": 30, "B": 15, "C": 10, "D": 24, "E": 10, "F": 16, "G": 18}.items():
        ws.column_dimensions[col].width = w

    # trees
    ws = wb.create_sheet("Trees")
    trees = result.get("trees", [])
    ws.append([f"Decision trees (genotype-by-genotype) — {len(trees)} option(s), best first"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Order", "Root", "Evidence", "Weakest link (%)", "Analyses", "Markers"])
    _hdr(ws)
    for i, t in enumerate(trees, 1):
        mk = sorted(f"{c[0].split('_chr')[-1]}:{c[1]}" for c in t["markers"])
        ws.append([i, _loc(t["node"]), "STRONG" if t["forte"] else "uses N",
                   t["wmin"], t["depth"], ", ".join(mk)])
        for c in ws[ws.max_row]:
            c.border = _BD
    for col, w in {"A": 7, "B": 12, "C": 14, "D": 16, "E": 10, "F": 36}.items():
        ws.column_dimensions[col].width = w
    ws.append([])

    def flow(t, idx):
        m = t["node"]
        ws.append([f"TREE {idx} — root {_loc(m)} (indel {m['indel']} bp):"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="1F4E78")
        for s, sub in t["groups"].items():
            g = sub[0]
            ws.append([f"  ROOT = {HRMLAB[s]} → " + ", ".join(_name(names, a) for a in g)])
            if len(g) > 1:
                dr = sub[1]["node"]
                det = ", ".join(_name(names, a) + "=" + HRMLAB[dr["labels"][a]] for a in g)
                ws.append([f"       2nd analysis {_loc(dr)} (indel {dr['indel']} bp, "
                           f"sep {sub[1].get('nodesep', '?')}%): {det}"])

    for i, t in enumerate(trees, 1):
        flow(t, i)
        ws.append([])
    ws.column_dimensions["A"].width = 100

    # primers
    ws = wb.create_sheet("Primers")
    ws.append(["Marker", "Indel (bp)", "Genotypes", "Forward primer", "Tm F",
               "Reverse primer", "Tm R", "Amplicon L/S (bp)", "Mapping hits",
               "In-silico PCR products"])
    _hdr(ws)
    used = {}
    for t in trees:
        used.update(_tree_markers(t))
    if not used:
        used = {(r["chrom"], r["pos"]): r for r in result.get("usable", [])}
    for key, r in sorted(used.items(), key=lambda kv: -kv[1]["indel"]):
        p = r["primers"]["pairs"][0]
        gh = r.get("genome_hits")
        gh = ", ".join(f"{k}:{v}" for k, v in gh.items()) if isinstance(gh, dict) else str(gh)
        ws.append([_loc(r), r["indel"], "/".join(r["labels"].get(a, "·") for a in samples),
                   p["left_seq"], p["left_tm"], p["right_seq"], p["right_tm"],
                   f"{len(p['amplicon_long'])}/{len(p['amplicon_short'])}", gh, _insilico_str(r)])
        for c in ws[ws.max_row]:
            c.border = _BD
    for col, w in {"A": 11, "B": 9, "C": 24, "D": 26, "E": 7, "F": 26, "G": 7,
                   "H": 16, "I": 16, "J": 22}.items():
        ws.column_dimensions[col].width = w

    # method
    ws = wb.create_sheet("Method")
    for line in [
        "DArTseq_InDel_Finder — genotype-by-genotype analysis",
        "",
        "Separation is computed between the predicted curve of each sample, according to its genotype:",
        "  A = long homoduplex | B = short homoduplex | AB = heterozygote (homo + heteroduplex).",
        "Two samples with the same genotype have the same curve and do NOT separate.",
        "N (double null = both tags absent) has an UNKNOWN sequence: no curve is predicted for it,",
        "and pairs that rely on N are reported as 'uses N (confirm)', never as STRONG.",
        "",
        "IMPORTANT: the difference-plot peak (%) is an UNCALIBRATED, RELATIVE prediction from a",
        "local melting model. It is NOT an instrument-calibrated value nor a probability of assay",
        "success. Use it to rank and compare candidates, and confirm the absolute melting behavior",
        "with uMelt (amplicons.fasta) and, ultimately, on your own HRM instrument.",
        "",
        f"Class by predicted peak (relative): >={cfg.hrm_easy:.0f}% easy | "
        f">={cfg.hrm_likely:.0f}% likely | below that, difficult. (Thresholds are generic heuristics.)",
        f"Prediction conditions: Mg2+ {cfg.melt_mg} mM, Mono+ {cfg.melt_mono} mM "
        "(Zimm-Bragg/Ising model, NN SantaLucia 1998; heterozygote via het_curve).",
        "",
        "In-silico PCR: exact-match screen against the reference, products counted per haplotype",
        "(homeolog amplifying once per haplotype is expected; >1 in one haplotype flags non-specific).",
        "",
        "Reminder: DArTseq does not report indels; they are INFERRED from co-located tags.",
        "Bench validation is required (indels inferred from presence/absence).",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 115

    xlsx = out_dir / "HRM_markers.xlsx"
    wb.save(xlsx)

    # amplicons for uMelt
    fasta = out_dir / "amplicons.fasta"
    with open(fasta, "w") as f:
        for key, r in used.items():
            p = r["primers"]["pairs"][0]
            tag = f"{_loc(r)}_indel{r['indel']}"
            f.write(f">{tag}_long\n{p['amplicon_long']}\n>{tag}_short\n{p['amplicon_short']}\n")

    img_dir = _save_images(result, cfg, names, out_dir / "curves")
    tsv_json = _write_machine_readable(result, cfg, out_dir, names, used)
    out = dict(xlsx=xlsx, fasta=fasta, images=img_dir)
    out.update(tsv_json)
    return out


def _gstr(r, samples):
    return "/".join(r["labels"].get(a, "·") for a in samples)


def _write_machine_readable(result, cfg, out_dir, names, used):
    """Write machine-readable TSV + JSON alongside the XLSX (reproducible, pipeline-friendly)."""
    out_dir = Path(out_dir)
    samples = result["samples"]

    # candidates.tsv — every characterized InDel candidate
    cand_tsv = out_dir / "candidates.tsv"
    with open(cand_tsv, "w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["chrom", "pos", "idA", "idB", "src", "indel_bp", "identity",
                    "exact_anchor", "homopolymer", "sep_clean", "n_clean", "incons_n",
                    "genotypes(" + "/".join(samples) + ")"])
        for o in result.get("out", []):
            w.writerow([o["chrom"], o["pos"], o["idA"], o["idB"], o["src"], o["indel"],
                        o["ident"], o.get("exact_anchor", 0), int(o["hp"]),
                        o["sep_clean"], o["n_clean"], o["incons_n"], _gstr(o, samples)])

    # primers.tsv — designed markers with primers + in-silico PCR
    prim_tsv = out_dir / "primers.tsv"
    with open(prim_tsv, "w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["marker", "indel_bp", "genotypes", "forward", "tm_f", "reverse", "tm_r",
                    "amplicon_long", "amplicon_short", "mapping_hits",
                    "insilico_products", "insilico_specific"])
        for key, r in sorted(used.items(), key=lambda kv: -kv[1]["indel"]):
            p = r["primers"]["pairs"][0]
            gh = r.get("genome_hits")
            gh = ";".join(f"{k}:{v}" for k, v in gh.items()) if isinstance(gh, dict) else str(gh)
            ins = r.get("insilico", {})
            per = ";".join(f"{k}:{v}" for k, v in ins.get("per_hap", {}).items())
            w.writerow([_loc(r), r["indel"], _gstr(r, samples), p["left_seq"], p["left_tm"],
                        p["right_seq"], p["right_tm"], len(p["amplicon_long"]),
                        len(p["amplicon_short"]), gh, per, int(ins.get("specific", False))])

    # per_pair.tsv — best markers per sample pair
    pp_tsv = out_dir / "per_pair.tsv"
    with open(pp_tsv, "w", newline="") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["sample_x", "sample_y", "marker", "indel_bp",
                    "predicted_peak_pct_uncalibrated", "class", "genotype_x", "genotype_y",
                    "evidence"])
        for (x, y) in itertools.combinations(samples, 2):
            for d in result.get("per_pair", {}).get((x, y), [])[:3]:
                w.writerow([_name(names, x), _name(names, y), _loc(d["marker"]),
                            d["marker"]["indel"], d["peak"], _klass(cfg, d["peak"]),
                            d["sx"], d["sy"], "STRONG" if d["forte"] else "uses N"])

    # run_summary.json — parameters, counts, chosen tree
    def tree_dict(t):
        return dict(root=_loc(t["node"]), root_indel_bp=t["node"]["indel"],
                    evidence="STRONG" if t["forte"] else "uses N",
                    weakest_link_pct=t["wmin"], analyses=t["depth"],
                    markers=sorted(f"{c[0].split('_chr')[-1]}:{c[1]}" for c in t["markers"]))
    summary = dict(
        tool="DArTseq_InDel_Finder",
        note=("Predicted difference-plot peaks are uncalibrated, relative values from a local "
              "melting model — not instrument-calibrated and not a probability of assay success. "
              "Indels are inferred from DArTseq presence/absence and require bench validation."),
        mode=result["mode"],
        samples=[{"code": s, "name": _name(names, s)} for s in samples],
        parameters=dict(
            melt_mg_mM=cfg.melt_mg, melt_mono_mM=cfg.melt_mono,
            indel_min=cfg.indel_min, indel_max=cfg.indel_max, flank=cfg.flank,
            min_ident=cfg.min_ident, min_anchor=cfg.min_anchor,
            min_exact_anchor=cfg.min_exact_anchor,
            hrm_easy_pct=cfg.hrm_easy, hrm_likely_pct=cfg.hrm_likely,
            insilico_filter=cfg.insilico_filter),
        counts=dict(candidates=len(result.get("out", [])),
                    usable_markers=len(result.get("usable", []))),
        trees=[tree_dict(t) for t in result.get("trees", [])],
    )
    if result["mode"] == "panel":
        mp = result.get("panel", {})
        summary["panel"] = dict(
            markers=[_loc(o) for o, _ in mp.get("panel", [])],
            n_pairs=mp.get("n_pairs"), resolved=mp.get("resolved"),
            resolvable=mp.get("resolvable"),
            unresolved=[[a, b] for a, b in mp.get("unresolved", [])])
    json_path = out_dir / "run_summary.json"
    with open(json_path, "w") as fh:
        json.dump(summary, fh, indent=2)

    return dict(candidates_tsv=cand_tsv, primers_tsv=prim_tsv,
                per_pair_tsv=pp_tsv, summary_json=json_path)
