"""Saídas: planilha XLSX (separação por par, árvores, primers) e imagens por amostra.

Genérico: usa um mapa opcional código->nome (metadata CSV com colunas 'code,name');
sem ele, usa os próprios códigos das amostras.
"""
import csv
import itertools
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .genotype import grid_of, pair_separation

HRMLAB = {"A": "LONGO", "B": "CURTO", "AB": "HETEROZIGOTO", "N": "N (presença/ausência)"}
_HF = PatternFill("solid", fgColor="1F4E78")
_HFONT = Font(color="FFFFFF", bold=True)
_THIN = Side(style="thin", color="CCCCCC")
_BD = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_PALETTE = ["#1f4e78", "#c55a11", "#2e7d32", "#7030a0", "#b00020", "#008b8b",
            "#8b4513", "#444444", "#d4a017", "#1565c0"]


def load_names(metadata_csv):
    names = {}
    if metadata_csv and Path(metadata_csv).exists():
        with open(metadata_csv, newline="") as fh:
            for row in csv.DictReader(fh):
                key = (row.get("code") or row.get("codigo") or "").strip()
                val = (row.get("name") or row.get("nome") or "").strip()
                if key:
                    names[key] = val or key
    return names


def _name(names, code):
    return names.get(code, code)


def _loc(r):
    return f"{r['chrom'].split('_chr')[-1]}:{r['pos']}"


def _hdr(ws):
    for c in ws[ws.max_row]:
        c.fill = _HF
        c.font = _HFONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _tree_markers(t):
    keys = {}

    def rec(entry):
        _, st = entry
        if st.get("leaf"):
            return
        keys[(st["node"]["chrom"], st["node"]["pos"])] = st["node"]
        for _, sub in st["groups"].items():
            rec(sub)

    keys[(t["node"]["chrom"], t["node"]["pos"])] = t["node"]
    for _, sub in t["groups"].items():
        rec(sub)
    return keys


def _save_images(result, cfg, names, img_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    img_dir = Path(img_dir)
    img_dir.mkdir(parents=True, exist_ok=True)
    grid = grid_of(cfg)
    samples = result["samples"]
    color = {s: _PALETTE[i % len(_PALETTE)] for i, s in enumerate(samples)}

    used = {}
    for t in result.get("trees", []):
        used.update(_tree_markers(t))
    if not used and result.get("usable"):  # painel: usa os do painel
        used = {(r["chrom"], r["pos"]): r for r in result["usable"][:12]}

    for key, r in used.items():
        fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4))
        ref = None
        for s in samples:
            st = r["labels"].get(s, "miss")
            if st == "miss":
                continue
            c = r["curves"][st]
            if ref is None:
                ref, refname = c, _name(names, s)
            a1.plot(grid, c * 100, color=color[s], label=f"{_name(names, s)} [{HRMLAB[st]}]")
            a2.plot(grid, (c - ref) * 100, color=color[s])
        a1.set_xlabel("Temperatura (°C)")
        a1.set_ylabel("Fluoresc. normalizada (%)")
        a1.set_title("Curvas por amostra (genótipo real)")
        a1.legend(fontsize=7)
        a1.grid(alpha=.3)
        a2.axhline(0, color="k", lw=.6)
        a2.set_xlabel("Temperatura (°C)")
        a2.set_ylabel(f"Diferença vs {refname} (%)")
        a2.set_title("Difference plot")
        a2.grid(alpha=.3)
        fig.suptitle(f"{_loc(r)} indel {r['indel']} pb — separa as amostras pelo genótipo")
        fig.tight_layout()
        fig.savefig(img_dir / f"{_loc(r).replace(':', '_')}_indel{r['indel']}.png", dpi=110)
        plt.close(fig)
    return img_dir


def _klass(cfg, peak):
    return "Fácil" if peak >= cfg.hrm_easy else ("Provável" if peak >= cfg.hrm_likely else "Difícil")


def write_report(result, cfg, out_dir, names=None):
    names = names or {}
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    samples = result["samples"]
    wb = Workbook()

    # separação por par
    ws = wb.active
    ws.title = "Separacao_por_par"
    ws.append(["Separação entre amostras no HRM — por par, usando o GENÓTIPO real de cada amostra"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Par", "Melhor marcador", "Indel (pb)", "Pico difference plot (%)",
               "Classe", "Genótipos (X/Y)", "Evidência"])
    _hdr(ws)
    per_pair = result.get("per_pair", {})
    for (x, y) in itertools.combinations(samples, 2):
        rows = per_pair.get((x, y), [])
        lab = f"{_name(names, x)} × {_name(names, y)}"
        if not rows:
            ws.append([lab, "— nenhum —", "", "", "", "", ""])
            continue
        for d in rows[:3]:
            ws.append([lab, _loc(d["marker"]), d["marker"]["indel"], d["peak"],
                       _klass(cfg, d["peak"]), f"{d['sx']}/{d['sy']}",
                       "FORTE" if d["forte"] else "usa N (confirmar)"])
            for c in ws[ws.max_row]:
                c.border = _BD
        ws.append([])
    for col, w in {"A": 30, "B": 15, "C": 10, "D": 22, "E": 10, "F": 14, "G": 18}.items():
        ws.column_dimensions[col].width = w

    # árvores
    ws = wb.create_sheet("Arvores")
    trees = result.get("trees", [])
    ws.append([f"Árvores de decisão (genótipo-a-genótipo) — {len(trees)} opção(ões), melhor primeiro"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(["Ordem", "Raiz", "Evidência", "Elo+fraco (%)", "Análises", "Marcadores"])
    _hdr(ws)
    for i, t in enumerate(trees, 1):
        mk = sorted(f"{c[0].split('_chr')[-1]}:{c[1]}" for c in t["markers"])
        ws.append([i, _loc(t["node"]), "FORTE" if t["forte"] else "usa N",
                   t["wmin"], t["depth"], ", ".join(mk)])
        for c in ws[ws.max_row]:
            c.border = _BD
    for col, w in {"A": 7, "B": 12, "C": 14, "D": 14, "E": 10, "F": 36}.items():
        ws.column_dimensions[col].width = w
    ws.append([])

    def flow(t, idx):
        m = t["node"]
        ws.append([f"ÁRVORE {idx} — raiz {_loc(m)} (indel {m['indel']} pb):"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="1F4E78")
        for s, sub in t["groups"].items():
            g = sub[0]
            ws.append([f"  RAIZ = {HRMLAB[s]} → " + ", ".join(_name(names, a) for a in g)])
            if len(g) > 1:
                dr = sub[1]["node"]
                det = ", ".join(_name(names, a) + "=" + HRMLAB[dr["labels"][a]] for a in g)
                ws.append([f"       2ª análise {_loc(dr)} (indel {dr['indel']} pb, "
                           f"sep {sub[1].get('nodesep', '?')}%): {det}"])

    for i, t in enumerate(trees, 1):
        flow(t, i)
        ws.append([])
    ws.column_dimensions["A"].width = 100

    # primers
    ws = wb.create_sheet("Primers")
    ws.append(["Marcador", "Indel (pb)", "Genótipos", "Primer Forward", "Tm F",
               "Primer Reverse", "Tm R", "Amplicon L/C (pb)", "Especificidade"])
    _hdr(ws)
    used = {}
    for t in trees:
        used.update(_tree_markers(t))
    if not used:
        used = {(r["chrom"], r["pos"]): r for r in result.get("usable", [])}
    for key, r in sorted(used.items(), key=lambda kv: -kv[1]["indel"]):
        p = r["primers"]["pairs"][0]
        gh = r.get("genome_hits")
        gh = ", ".join(f"{k}:{v}" for k, v in gh.items()) if isinstance(gh, dict) else str(gh)
        ws.append([_loc(r), r["indel"], "/".join(r["labels"].get(a, "·") for a in samples),
                   p["left_seq"], p["left_tm"], p["right_seq"], p["right_tm"],
                   f"{len(p['amplicon_long'])}/{len(p['amplicon_short'])}", gh])
        for c in ws[ws.max_row]:
            c.border = _BD
    for col, w in {"A": 11, "B": 9, "C": 24, "D": 26, "E": 7, "F": 26, "G": 7, "H": 14, "I": 16}.items():
        ws.column_dimensions[col].width = w

    # método
    ws = wb.create_sheet("Metodo")
    for line in [
        "indel-hrm-finder — análise genótipo-a-genótipo",
        "",
        "A separação é calculada entre a curva REAL de cada amostra, conforme o genótipo:",
        "  A = homoduplex longo | B = homoduplex curto | AB = heterozigoto (homo + heteroduplex).",
        "Duas amostras com o mesmo genótipo têm a mesma curva e NÃO se separam.",
        "",
        f"Classe por pico no difference plot: ≥{cfg.hrm_easy:.0f}% fácil | "
        f"≥{cfg.hrm_likely:.0f}% provável | abaixo difícil.",
        f"Condições da predição: Mg²⁺ {cfg.melt_mg} mM, Mono⁺ {cfg.melt_mono} mM "
        "(modelo Zimm-Bragg/Ising, NN SantaLucia 1998; heterozigoto via het_curve).",
        "",
        "Lembrete: o DArTseq não reporta indels; eles são INFERIDOS de tags co-localizadas.",
        "Validação de bancada é necessária (indels inferidos de presença/ausência).",
    ]:
        ws.append([line])
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 115

    xlsx = out_dir / "HRM_marcadores.xlsx"
    wb.save(xlsx)

    # amplicons para uMelt
    fasta = out_dir / "amplicons.fasta"
    with open(fasta, "w") as f:
        for key, r in used.items():
            p = r["primers"]["pairs"][0]
            tag = f"{_loc(r)}_indel{r['indel']}"
            f.write(f">{tag}_longo\n{p['amplicon_long']}\n>{tag}_curto\n{p['amplicon_short']}\n")

    img_dir = _save_images(result, cfg, names, out_dir / "curvas")
    return dict(xlsx=xlsx, fasta=fasta, images=img_dir)
